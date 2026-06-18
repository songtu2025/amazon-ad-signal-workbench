import hashlib
import json
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_PARENT_ASIN_SNAPSHOT_ROOT = PROJECT_ROOT / "data" / "parent_asin_snapshots"
DEFAULT_SOURCE_SHEET = "按日导出"
SOURCE_TYPE = "销售表现"
SOURCE_TABLE = "parent_asin_sales_performance"
PARENT_ASIN_MAP_FILE = "parent_asin_product_map.json"


@dataclass(frozen=True)
class ParentAsinImportResult:
    source_file: Path
    source_sheet: str
    file_hash: str
    marketplace_id: int | None
    marketplace_code: str
    start_date: str
    end_date: str
    imported_at: str
    row_count: int
    skipped_count: int
    duplicate_count: int
    records: list[dict[str, Any]]


def import_parent_asin_workbook(
    *,
    source_file: Path,
    marketplace_code: str,
    marketplace_id: int | None = None,
    source_sheet: str = DEFAULT_SOURCE_SHEET,
) -> ParentAsinImportResult:
    source_file = Path(source_file)
    _validate_import_request(source_file, marketplace_code)

    workbook = load_workbook(source_file, read_only=True, data_only=True)
    try:
        if source_sheet not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{source_sheet}")
        sheet = workbook[source_sheet]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("销售表现-父ASIN 文件为空")
        header_map = _header_map(header)
        _validate_headers(header_map)

        grouped_records: dict[tuple[str, str, str], dict[str, Any]] = {}
        skipped_count = 0
        duplicate_count = 0
        observed_dates: list[str] = []
        for source_row_number, row in enumerate(rows, start=2):
            record = _build_record(
                row,
                header_map,
                source_file=source_file,
                source_sheet=source_sheet,
                source_row_number=source_row_number,
                marketplace_id=marketplace_id,
                marketplace_code=marketplace_code,
            )
            if record is None:
                skipped_count += 1
                continue
            observed_dates.append(record["date"])
            key = (record["marketplace_code"], record["parent_asin"], record["asin"])
            existing = grouped_records.get(key)
            if existing is None:
                grouped_records[key] = record
                continue
            duplicate_count += 1
            _merge_record(existing, record)

        records = sorted(grouped_records.values(), key=lambda item: (item["parent_asin"], item["asin"]))
        if not records:
            raise ValueError("销售表现-父ASIN 文件没有可导入的父子 ASIN 行")

        start_date = min(observed_dates)
        end_date = max(observed_dates)
        imported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        for record in records:
            record["start_date"] = start_date
            record["end_date"] = end_date
            record.pop("date", None)

        return ParentAsinImportResult(
            source_file=source_file,
            source_sheet=source_sheet,
            file_hash=_file_hash(source_file),
            marketplace_id=marketplace_id,
            marketplace_code=marketplace_code,
            start_date=start_date,
            end_date=end_date,
            imported_at=imported_at,
            row_count=len(records),
            skipped_count=skipped_count,
            duplicate_count=duplicate_count,
            records=records,
        )
    finally:
        workbook.close()


def write_parent_asin_snapshot(
    result: ParentAsinImportResult,
    *,
    snapshot_root: Path = DEFAULT_PARENT_ASIN_SNAPSHOT_ROOT,
) -> Path:
    snapshot_id = f"parent_asin_{result.marketplace_code}_{result.start_date}_{result.end_date}_{result.file_hash[:8]}"
    safe_snapshot_id = snapshot_id.replace("-", "")
    snapshot_dir = Path(snapshot_root) / safe_snapshot_id
    normalized_dir = snapshot_dir / "normalized"
    normalized_dir.mkdir(parents=True, exist_ok=True)

    manifest = {
        "snapshot_id": safe_snapshot_id,
        "source": "erp_export",
        "source_type": SOURCE_TYPE,
        "source_table": SOURCE_TABLE,
        "source_file": str(result.source_file),
        "source_sheet": result.source_sheet,
        "file_hash": result.file_hash,
        "marketplace_id": result.marketplace_id,
        "marketplace_code": result.marketplace_code,
        "start_date": result.start_date,
        "end_date": result.end_date,
        "imported_at": result.imported_at,
        "row_count": result.row_count,
        "skipped_count": result.skipped_count,
        "duplicate_count": result.duplicate_count,
        "status": "success",
    }
    _write_json(snapshot_dir / "manifest.json", manifest)
    _write_json(normalized_dir / PARENT_ASIN_MAP_FILE, result.records)
    return snapshot_dir


def _validate_import_request(source_file: Path, marketplace_code: str) -> None:
    if not source_file.exists():
        raise FileNotFoundError(f"销售表现-父ASIN 文件不存在：{source_file}")
    if source_file.suffix.lower() != ".xlsx":
        raise ValueError("销售表现-父ASIN 导入文件必须是 .xlsx")
    if not marketplace_code:
        raise ValueError("marketplace_code 必填")


def _validate_headers(header_map: dict[str, int]) -> None:
    required_columns = ["日期", "父ASIN", "ASIN", "店铺/站点", "国家", "产品名称", "SKU"]
    missing = [column for column in required_columns if column not in header_map]
    if missing:
        raise ValueError(f"销售表现-父ASIN 文件缺少必要列：{', '.join(missing)}")


def _header_map(header: tuple[Any, ...]) -> dict[str, int]:
    return {str(value).strip(): index for index, value in enumerate(header) if value not in (None, "")}


def _build_record(
    row: tuple[Any, ...],
    header_map: dict[str, int],
    *,
    source_file: Path,
    source_sheet: str,
    source_row_number: int,
    marketplace_id: int | None,
    marketplace_code: str,
) -> dict[str, Any] | None:
    row_date = _date_string(_cell(row, header_map, "日期"))
    parent_asin = _string(_cell(row, header_map, "父ASIN"))
    asin = _string(_cell(row, header_map, "ASIN"))
    if not row_date or not parent_asin or not asin:
        return None

    return {
        "marketplace_id": marketplace_id,
        "marketplace_code": marketplace_code,
        "date": row_date,
        "parent_asin": parent_asin,
        "asin": asin,
        "sku": _string(_cell(row, header_map, "SKU")),
        "msku": _string(_cell(row, header_map, "MSKU")),
        "product_name": _string(_cell(row, header_map, "产品名称")),
        "parent_product_name": _string(_cell(row, header_map, "父产品名称")),
        "spu": _string(_cell(row, header_map, "SPU")),
        "shop_site": _string(_cell(row, header_map, "店铺/站点")),
        "country": _string(_cell(row, header_map, "国家")),
        "orders": _integer(_cell(row, header_map, "订单量")) or 0,
        "sales": _number(_cell(row, header_map, "销售额")) or 0,
        "ad_spend": _absolute_number(_cell(row, header_map, "广告花费")) or 0,
        "ad_sales": _number(_cell(row, header_map, "广告销售额")) or 0,
        "ad_orders": _integer(_cell(row, header_map, "广告订单量")) or 0,
        "organic_orders": _integer(_cell(row, header_map, "自然订单量")) or 0,
        "sessions_total": _integer(_cell(row, header_map, "Sessions-Total")),
        "cvr": _number(_cell(row, header_map, "CVR")),
        "source_type": SOURCE_TYPE,
        "source_table": SOURCE_TABLE,
        "source_file": str(source_file),
        "source_sheet": source_sheet,
        "source_row_number": source_row_number,
        "source_record_id": f"parent-asin:{marketplace_code}:{parent_asin}:{asin}",
        "import_batch_id": None,
    }


def _merge_record(existing: dict[str, Any], record: dict[str, Any]) -> None:
    for field in ["orders", "sales", "ad_spend", "ad_sales", "ad_orders", "organic_orders"]:
        existing[field] = round((_number(existing.get(field)) or 0) + (_number(record.get(field)) or 0), 4)
    for field in ["sku", "msku", "product_name", "parent_product_name", "spu", "shop_site", "country"]:
        if not existing.get(field) and record.get(field):
            existing[field] = record[field]


def _cell(row: tuple[Any, ...], header_map: dict[str, int], column: str) -> Any:
    index = header_map.get(column)
    if index is None or index >= len(row):
        return None
    return row[index]


def _string(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _date_string(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()[:10]


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    try:
        return int(float(str(value).replace(",", "").replace("%", "")))
    except (TypeError, ValueError):
        return None


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(str(value).replace(",", "").replace("%", ""))
    except (TypeError, ValueError):
        return None


def _absolute_number(value: Any) -> float | None:
    number = _number(value)
    if number is None:
        return None
    return abs(number)


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
