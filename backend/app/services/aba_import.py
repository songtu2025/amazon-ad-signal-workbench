import hashlib
import json
import re
from dataclasses import dataclass
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[3]
DEFAULT_ABA_SNAPSHOT_ROOT = PROJECT_ROOT / "data" / "aba_snapshots"
DEFAULT_SOURCE_SHEET = "关键词搜索热度"
SOURCE_TYPE = "ABA导出"
SOURCE_TABLE = "aba_search_term_snapshots"
TOP_RANK_LIMIT = 1000
TOP_RANK_INDEX_FILE = "aba_search_term_top1000.json"


@dataclass(frozen=True)
class AbaImportResult:
    source_file: Path
    source_sheet: str
    file_hash: str
    marketplace_id: int | None
    marketplace_code: str
    start_date: str
    end_date: str
    imported_at: str
    row_count: int
    skipped_count: int
    duplicate_count: int
    records: list[dict[str, Any]]


def normalize_query(value: Any) -> str:
    if value in (None, ""):
        return ""
    return re.sub(r"\s+", " ", str(value).strip().lower())


def import_aba_workbook(
    *,
    source_file: Path,
    marketplace_code: str,
    start_date: str,
    end_date: str,
    marketplace_id: int | None = None,
    source_sheet: str = DEFAULT_SOURCE_SHEET,
) -> AbaImportResult:
    source_file = Path(source_file)
    _validate_import_request(source_file, marketplace_code, start_date, end_date)

    workbook = load_workbook(source_file, read_only=True, data_only=True)
    try:
        if source_sheet not in workbook.sheetnames:
            raise ValueError(f"工作表不存在：{source_sheet}")
        sheet = workbook[source_sheet]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            raise ValueError("ABA 文件为空")
        header_map = _header_map(header)
        _validate_headers(header_map)

        deduped_records: dict[str, dict[str, Any]] = {}
        skipped_count = 0
        duplicate_count = 0
        for source_row_number, row in enumerate(rows, start=2):
            record = _build_record(
                row,
                header_map,
                source_file=source_file,
                source_sheet=source_sheet,
                source_row_number=source_row_number,
                marketplace_id=marketplace_id,
                marketplace_code=marketplace_code,
                start_date=start_date,
                end_date=end_date,
            )
            if record is None:
                skipped_count += 1
                continue
            normalized_query = record["normalized_query"]
            existing = deduped_records.get(normalized_query)
            if existing is not None:
                duplicate_count += 1
                if record["search_frequency_rank"] < existing["search_frequency_rank"]:
                    deduped_records[normalized_query] = record
                continue
            deduped_records[normalized_query] = record

        records = sorted(deduped_records.values(), key=lambda item: item["search_frequency_rank"])
        imported_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
        return AbaImportResult(
            source_file=source_file,
            source_sheet=source_sheet,
            file_hash=_file_hash(source_file),
            marketplace_id=marketplace_id,
            marketplace_code=marketplace_code,
            start_date=start_date,
            end_date=end_date,
            imported_at=imported_at,
            row_count=len(records),
            skipped_count=skipped_count,
            duplicate_count=duplicate_count,
            records=records,
        )
    finally:
        workbook.close()


def write_aba_snapshot(
    result: AbaImportResult,
    *,
    snapshot_root: Path = DEFAULT_ABA_SNAPSHOT_ROOT,
) -> Path:
    snapshot_id = f"aba_{result.marketplace_code}_{result.start_date}_{result.end_date}_{result.file_hash[:8]}"
    safe_snapshot_id = snapshot_id.replace("-", "")
    snapshot_dir = Path(snapshot_root) / safe_snapshot_id
    normalized_dir = snapshot_dir / "normalized"
    normalized_dir.mkdir(parents=True, exist_ok=True)

    manifest = {
        "snapshot_id": safe_snapshot_id,
        "source": "erp_export",
        "source_type": SOURCE_TYPE,
        "source_file": str(result.source_file),
        "source_sheet": result.source_sheet,
        "file_hash": result.file_hash,
        "marketplace_id": result.marketplace_id,
        "marketplace_code": result.marketplace_code,
        "start_date": result.start_date,
        "end_date": result.end_date,
        "imported_at": result.imported_at,
        "row_count": result.row_count,
        "top_rank_limit": TOP_RANK_LIMIT,
        "top_rank_index_count": len(_top_rank_records(result.records)),
        "skipped_count": result.skipped_count,
        "duplicate_count": result.duplicate_count,
        "status": "success",
    }
    _write_json(snapshot_dir / "manifest.json", manifest)
    _write_json(normalized_dir / f"{SOURCE_TABLE}.json", result.records)
    _write_json(normalized_dir / TOP_RANK_INDEX_FILE, _top_rank_records(result.records))
    return snapshot_dir


def _top_rank_records(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    return [record for record in records if record["search_frequency_rank"] <= TOP_RANK_LIMIT]


def _validate_import_request(source_file: Path, marketplace_code: str, start_date: str, end_date: str) -> None:
    if not source_file.exists():
        raise FileNotFoundError(f"ABA 文件不存在：{source_file}")
    if source_file.suffix.lower() != ".xlsx":
        raise ValueError("ABA 导入文件必须是 .xlsx")
    if not marketplace_code:
        raise ValueError("marketplace_code 必填")
    if start_date > end_date:
        raise ValueError("start_date 不能晚于 end_date")


def _validate_headers(header_map: dict[str, int]) -> None:
    missing = [column for column in ["关键词", "搜索量排名"] if column not in header_map]
    if missing:
        raise ValueError(f"ABA 文件缺少必要列：{', '.join(missing)}")
    if "#1 ASIN" not in header_map:
        raise ValueError("ABA 文件至少需要包含一组前三 ASIN 字段")


def _header_map(header: tuple[Any, ...]) -> dict[str, int]:
    return {str(value).strip(): index for index, value in enumerate(header) if value not in (None, "")}


def _build_record(
    row: tuple[Any, ...],
    header_map: dict[str, int],
    *,
    source_file: Path,
    source_sheet: str,
    source_row_number: int,
    marketplace_id: int | None,
    marketplace_code: str,
    start_date: str,
    end_date: str,
) -> dict[str, Any] | None:
    search_term = _string(_cell(row, header_map, "关键词"))
    normalized_query = normalize_query(search_term)
    search_frequency_rank = _integer(_cell(row, header_map, "搜索量排名"))
    if not normalized_query or search_frequency_rank is None:
        return None

    record = {
        "marketplace_id": marketplace_id,
        "marketplace_code": marketplace_code,
        "start_date": start_date,
        "end_date": end_date,
        "search_term": search_term,
        "normalized_query": normalized_query,
        "search_frequency_rank": search_frequency_rank,
        "rank_change_type": _string(_cell(row, header_map, "排名变化")),
        "rank_change_value": _integer(_cell(row, header_map, "变化名次")),
        "top3_click_share": _number(_cell(row, header_map, "前三ASIN点击份额（%）")),
        "top3_conversion_share": _number(_cell(row, header_map, "前三ASIN转化份额（%）")),
        "source_type": SOURCE_TYPE,
        "source_table": SOURCE_TABLE,
        "source_file": str(source_file),
        "source_sheet": source_sheet,
        "source_row_number": source_row_number,
        "import_batch_id": None,
    }
    for rank in (1, 2, 3):
        prefix = f"#{rank}"
        target_prefix = f"top{rank}"
        record[f"{target_prefix}_asin"] = _string(_cell(row, header_map, f"{prefix} ASIN"))
        record[f"{target_prefix}_title"] = _string(_cell(row, header_map, f"{prefix} ASIN标题"))
        record[f"{target_prefix}_click_share"] = _number(_cell(row, header_map, f"{prefix} ASIN点击份额（%）"))
        record[f"{target_prefix}_conversion_share"] = _number(_cell(row, header_map, f"{prefix} ASIN转化份额（%）"))
    record["source_record_id"] = (
        f"aba:{marketplace_code}:{start_date}:{end_date}:{record['normalized_query']}"
    )
    return record


def _cell(row: tuple[Any, ...], header_map: dict[str, int], column: str) -> Any:
    index = header_map.get(column)
    if index is None or index >= len(row):
        return None
    return row[index]


def _string(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).strip()


def _integer(value: Any) -> int | None:
    if value in (None, ""):
        return None
    return int(float(value))


def _number(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as file:
        for chunk in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
